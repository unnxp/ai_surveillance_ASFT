import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def main():
    doc = docx.Document()

    # Base styling
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'TH Sarabun PSK'
    font.size = Pt(16)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("AI-Based Smart Surveillance System")
    r_title.bold = True
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("รายละเอียดโปรเจกต์และการออกแบบสถาปัตยกรรมระบบกล้องวงจรปิดอัจฉริยะวิเคราะห์พฤติกรรมและประเมินความเสี่ยงด้วย AI")
    r_sub.font.size = Pt(16)
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    r_sub.italic = True

    doc.add_paragraph() # Spacer

    # Helper function for section headings
    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(20)
        r.font.color.rgb = RGBColor(0x00, 0x44, 0x88)
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(17)
        r.font.color.rgb = RGBColor(0x00, 0x66, 0x99)
        return h

    # Section 1
    add_heading_1("1. ภาพรวมโปรเจกต์ (Project Overview)")
    doc.add_paragraph(
        "โครงการนี้คือระบบกล้องวงจรปิดอัจฉริยะ (AI-Based Smart Surveillance System) "
        "ซึ่งถูกออกแบบมาเพื่อเฝ้าระวัง ตรวจจับ และแจ้งเตือนเหตุการณ์เสี่ยงภัยหรืออาชญากรรมจากกล้องวงจรปิด IP Camera แบบ Real-time "
        "โดยมีเป้าหมายในการยกระดับระบบ CCTV ดั้งเดิม ให้สามารถวิเคราะห์เหตุการณ์ได้โดยอัตโนมัติ ช่วยลดภาระของเจ้าหน้าที่ควบคุมหน้าจอ "
        "และลดความผิดพลาดจากมนุษย์ (Human Error)"
    )

    doc.add_paragraph(
        "ข้อจำกัดและการออกแบบฮาร์ดแวร์ (Single-Machine Constraint):\n"
        "ระบบถูกออกแบบโครงสร้างสถาปัตยกรรมแบบ Distributed System on Single-Machine เพื่อให้สามารถประมวลผลวิดีโอจากกล้องวงจรปิดแบบเรียลไทม์ "
        "บนฮาร์ดแวร์ระดับเครื่องคอมพิวเตอร์พกพา (เช่น MSI Laptop: CPU AMD Ryzen 5 7535HS, GPU NVIDIA RTX 3050 4GB VRAM, RAM 16GB, OS Windows) "
        "โดยแบ่งแยกหน้าที่ระหว่าง Edge Processing (ประมวลผลสกัดฟีเจอร์ที่ต้นทาง) และ Central AI Server (ศูนย์กลางการตัดสินใจและจัดการข้อมูล) อย่างเป็นระบบ"
    )

    # Section 2
    add_heading_1("2. สถาปัตยกรรมระบบและกระบวนการทำงาน (System Architecture & Data Flow)")
    doc.add_paragraph(
        "สถาปัตยกรรมของระบบแบ่งออกเป็น 2 ส่วนหลัก เพื่อการใช้ทรัพยากรและแบนด์วิดท์อย่างมีประสิทธิภาพสูงสุด:"
    )

    add_heading_2("2.1 Edge Processing Nodes (โหนดประมวลผลต้นทางประจำกล้อง)")
    p_edge = doc.add_paragraph()
    p_edge.paragraph_format.left_indent = Inches(0.2)
    p_edge.add_run("• การดึงภาพและตรวจจับบุคคล: ").bold = True
    p_edge.add_run("เชื่อมต่อกล้อง IP Camera (RTSP/USB/File) และตรวจจับพิกัดบุคคลด้วย YOLOv8 + ByteTrack (Kalman Filter Smoothing)\n")
    p_edge.add_run("• การจำแนกใบหน้า (Face Recognition): ").bold = True
    p_edge.add_run("ตรวจจับพิกัดใบหน้าด้วย SCRFD และสกัด Feature Vector 512 มิติด้วย ArcFace (รันด้วย ONNX Runtime บน CPU เพื่อประหยัด VRAM)\n")
    p_edge.add_run("• การจำแนกพฤติกรรมเสี่ยง (Action Classification): ").bold = True
    p_edge.add_run("วิเคราะห์การกระทำด้วยโมเดล EfficientNet-B0 + GRU (Periodic Inference ทุกๆ 1.5–2 วินาทีเพื่อลดภาระ GPU)\n")
    p_edge.add_run("• ตรรกะการประเมินความเสี่ยงและพื้นที่: ").bold = True
    p_edge.add_run("ตรวจสอบ Geofencing (3 จุด: เท้า, เข่า, เอว), Loitering (การเตร็ดเตร่), Speed (ความเร็ววิ่ง), อาวุธ (Knife proximity) และสร้าง Heatmap ความหนาแน่น 32x32\n")
    p_edge.add_run("• Edge JSON Protocol: ").bold = True
    p_edge.add_run("บีบอัดฟีเจอร์ทั้งหมดเป็นแพ็กเกจ JSON ขนาดเล็ก (< 4KB/sec) ส่งเข้า Central AI Server")

    add_heading_2("2.2 Central AI Server (ศูนย์กลางการตัดสินใจ)")
    p_central = doc.add_paragraph()
    p_central.paragraph_format.left_indent = Inches(0.2)
    p_central.add_run("• Face 1:N Search & Access Control: ").bold = True
    p_central.add_run("รับเวกเตอร์ใบหน้ามาเปรียบเทียบในคลังข้อมูลบุคคลเฝ้าระวัง (Watchlist) ด้วย FAISS Vector Store เพื่อระบุสิทธิ์การเข้าออกพื้นที่\n")
    p_central.add_run("• Risk Fusion & Alert Management: ").bold = True
    p_central.add_run("รวมสัญญาณความเสี่ยง บันทึกประวัติเหตุการณ์ลง CSV และบันทึกภาพ Snapshot หน้าจอเหตุการณ์เสี่ยง\n")
    p_central.add_run("• Web Dashboard & LLM Summarizer: ").bold = True
    p_central.add_run("ให้บริการหน้าต่าง Web Dashboard สำหรับดู Live Stream, Alert Feed และสรุปรายงานเหตุการณ์ภาษาธรรมชาติด้วย Local LLM (Ollama + gemma2:2b)")

    # Section 3
    add_heading_1("3. องค์ประกอบหลักของโปรเจกต์ (Core Modules Breakdown)")
    
    add_heading_2("3.1 โมดูล ai/ (ปัญญาประดิษฐ์สกัดฟีเจอร์)")
    p_ai = doc.add_paragraph()
    p_ai.paragraph_format.left_indent = Inches(0.2)
    p_ai.add_run("1. detector.py: ").bold = True
    p_ai.add_run("YOLOv8 Person Detection (ประมวลผลด้วย FP16 บน GPU)\n")
    p_ai.add_run("2. tracker.py: ").bold = True
    p_ai.add_run("ByteTrack + Kalman Filter สำหรับทำ Smoothed Trajectory และวัดความเร็วเคลื่อนที่\n")
    p_ai.add_run("3. face_recognition.py: ").bold = True
    p_ai.add_run("SCRFD (Face Detect) + ArcFace (512-dim Embedding) รันบน CPU ด้วย ONNX\n")
    p_ai.add_run("4. action_classifier.py: ").bold = True
    p_ai.add_run("EfficientNet-B0 + GRU จำแนกพฤติกรรม 3 คลาส (Normal, Property Crime, Personal Violence) พร้อม Confidence Filter 70%")

    add_heading_2("3.2 โมดูล logic/ (ตรรกะและการวิเคราะห์ความเสี่ยง)")
    p_logic = doc.add_paragraph()
    p_logic.paragraph_format.left_indent = Inches(0.2)
    p_logic.add_run("1. risk_assessment.py: ").bold = True
    p_logic.add_run("ประเมินเงื่อนไข Geofence Polygon, เวลา Off-hours (22:00-06:00), เวลาเตร็ดเตร่ และระยะห่างอาวุธ\n")
    p_logic.add_run("2. heatmap.py: ").bold = True
    p_logic.add_run("สร้าง Heatmap Grid ขนาด 32x32 และตรวจจับการบุกรุกโซนที่ไม่เคยมีการสัญจร (Cold Zone Path Anomaly)\n")
    p_logic.add_run("3. alert_manager.py: ").bold = True
    p_logic.add_run("จัดการระบบ Log CSV และจับภาพ Snapshot เมื่อเกิดเหตุการณ์ระดับ HIGH / CRITICAL\n")
    p_logic.add_run("4. edge_protocol.py: ").bold = True
    p_logic.add_run("จัดรูปแบบโครงสร้างข้อมูล JSON สำหรับการสื่อสารระหว่าง Edge และ Central Server")

    add_heading_2("3.3 โมดูล configs/ & tools/ (การตั้งค่าและเครื่องมือช่วยเหลือ)")
    p_cfg = doc.add_paragraph()
    p_cfg.paragraph_format.left_indent = Inches(0.2)
    p_cfg.add_run("1. behavior_config.yaml: ").bold = True
    p_cfg.add_run("ไฟล์ตั้งค่าพิกัด Geofence Polygon, Thresholds และการทำงานของระบบ\n")
    p_cfg.add_run("2. tools/draw_geofence.py: ").bold = True
    p_cfg.add_run("เครื่องมือวาดพิกัด Geofence บนภาพวิดีโอแบบ Interactive\n")
    p_cfg.add_run("3. tools/train_action_classifier_3cls.py: ").bold = True
    p_cfg.add_run("สคริปต์ฝึกสอนโมเดลวิเคราะห์พฤติกรรมเสี่ยง 3 คลาส")

    # Section 4
    add_heading_1("4. การบริหารจัดการทรัพยากร GPU/CPU (VRAM Budgeting)")
    doc.add_paragraph(
        "เนื่องจากระบบต้องทำงานบน GPU RTX 3050 (ขีดจำกัด 4GB VRAM) จึงได้จัดสรรงบประมาณหน่วยความจำดังนี้:"
    )

    # Table for VRAM Budget
    table_vram = doc.add_table(rows=1, cols=3)
    table_vram.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table_vram.rows[0].cells
    headers = ["โมดูล / ส่วนประกอบ (Component)", "อุปกรณ์ (Device)", "ประมาณการ VRAM (Est. VRAM)"]
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(hdr_cells[i], "004488")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)

    vram_data = [
        ("YOLOv8n FP16 × 2 กล้อง", "GPU (CUDA)", "~500 MB"),
        ("Action Classifier (Periodic Inference)", "GPU (CUDA)", "~300 MB"),
        ("SCRFD + ArcFace (ONNX Runtime)", "CPU", "0 MB (ใช้ RAM System)"),
        ("CUDA Context & Graphic Driver", "GPU", "~400 MB"),
        ("รวม VRAM ที่ใช้งานทั้งหมด (Total)", "GPU", "~1,200 MB (จากงบสูงสุด 4,000 MB) ✅")
    ]

    for comp, dev, vram in vram_data:
        row_cells = table_vram.add_row().cells
        row_cells[0].text = comp
        row_cells[1].text = dev
        row_cells[2].text = vram
        for j, cell in enumerate(row_cells):
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

    doc.add_paragraph() # Spacer

    # Section 5
    add_heading_1("5. สถานะการพัฒนาในปัจจุบัน (Current Development Status)")
    
    add_heading_2("5.1 ส่วนที่เสร็จสมบูรณ์แล้ว (Completed ✅)")
    completed_items = [
        "ระบบตรวจจับและติดตามบุคคลรายกล้องด้วย YOLOv8 + ByteTrack (แยกอิสระกล้อง 1 และ 2)",
        "ระบบสกัด Feature Vector ใบหน้า 512 มิติด้วย SCRFD + ArcFace (ONNX บน CPU)",
        "โมเดลจำแนกพฤติกรรมเสี่ยง 3 คลาส (EfficientNet-B0 + GRU) รันแบบ Periodic Inference",
        "ระบบ Geofence Multi-point check, Loitering, Speed detection และ Weapon proximity (ตรวจจับมีด)",
        "ระบบ Heatmap 32x32 และตรวจจับการบุกรุกโซนเย็น (Cold Zone Anomaly)",
        "ระบบ Edge JSON Exporter Protocol (edge_protocol.py)",
        "ระบบ Alert Manager (บันทึก CSV + จับภาพ Snapshot)"
    ]
    for item in completed_items:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.left_indent = Inches(0.2)
        p_c.add_run("✓ ").bold = True
        p_c.add_run(item)

    add_heading_2("5.2 ขั้นตอนการพัฒนาถัดไป (Roadmap / Next Steps ⏳)")
    roadmap_items = [
        "Phase 5 (Central AI Server): พัฒนา FastAPI Server, สถาปัตยกรรม FAISS Vector Search สำหรับค้นหาใบหน้า 1:N และเชื่อมโยง Local LLM สรุปรายงานเหตุการณ์",
        "Phase 6 (Web Dashboard): สร้างส่วนติดต่อผู้ใช้แบบ Web Application สำหรับดู Live Stream, Alert Feed และ Heatmap"
    ]
    for item in roadmap_items:
        p_r = doc.add_paragraph()
        p_r.paragraph_format.left_indent = Inches(0.2)
        p_r.add_run("→ ").bold = True
        p_r.add_run(item)

    # Section 6
    add_heading_1("6. จุดเด่นทางเทคนิคของระบบ (Key Technical Highlights)")
    highlights = [
        ("Edge Feature Exporter Architecture: ", "สกัดฟีเจอร์สำคัญที่ต้นทางแล้วส่งต่อเฉพาะข้อความ JSON ขนาดเล็ก ช่วยประหยัด Bandwidth เครือข่ายอย่างมหาศาล"),
        ("Periodic Action Inference: ", "ประมวลผลโมเดลพฤติกรรมเฉพาะทุกๆ ช่วงเวลา 1.5–2 วินาที ช่วยลดภาระ GPU ได้มากกว่า 75% โดยไม่สูญเสียความแม่นยำในการตรวจจับ"),
        ("Multi-Factor Risk Assessment: ", "ผสมผสานสัญญาณความเสี่ยงรอบด้าน (ตำแหน่ง + เวลา + ความเร็ว + การเตร็ดเตร่ + ใบหน้า + พฤติกรรม + อาวุธ) ทำให้ผลการแจ้งเตือนมีความแม่นยำสูง")
    ]
    for title, desc in highlights:
        p_h = doc.add_paragraph()
        p_h.paragraph_format.left_indent = Inches(0.2)
        p_h.add_run("• " + title).bold = True
        p_h.add_run(desc)

    # Save output
    output_path = "AI_Surveillance_Project_Explanation.docx"
    doc.save(output_path)
    print(f"Successfully generated docx document at: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    main()
