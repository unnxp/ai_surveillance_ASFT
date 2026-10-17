import docx
import os

def main():
    original_path = "6610110655_AI-Based Smart Surveillance System.docx"
    updated_path = "6610110655_AI-Based Smart Surveillance System_Updated.docx"
    
    if not os.path.exists(original_path):
        print(f"Error: Original file '{original_path}' not found.")
        return
        
    doc = docx.Document(original_path)
    
    # ── 1. การแก้ไขข้อความใน Paragraphs ──────────────────────────────
    print("Updating paragraphs...")
    
    replacements = {
        # 5. ความสำคัญและที่มา
        "การเฝ้าระวังแบบดั้งเดิมยังคงต้องอาศัยเจ้าหน้าที่จับตามองหน้าจอ": (
            "ในปัจจุบันระบบกล้องวงจรปิด (CCTV) ถูกติดตั้งอย่างแพร่หลายในสถานที่สำคัญต่าง ๆ อย่างไรก็ตาม "
            "การเฝ้าระวังแบบดั้งเดิมยังคงต้องอาศัยเจ้าหน้าที่จับตามองหน้าจอตลอดเวลา ซึ่งมีข้อจำกัดด้านความสามารถ"
            "ในการตรวจจับเหตุการณ์ที่เกิดขึ้นพร้อมกันหลายจุด และอาจเกิดความผิดพลาดจากมนุษย์ (Human Error) ได้ "
            "โครงงานนี้จึงมุ่งพัฒนาระบบเฝ้าระวังพื้นที่อัตโนมัติแบบประหยัดแบนด์วิดท์ โดยใช้สถาปัตยกรรมแบบกระจายศูนย์ "
            "(Distributed Edge-Central System) นำเทคโนโลยีปัญญาประดิษฐ์มาตรวจจับและจำแนกพฤติกรรมผิดปกติ 3 คลาสหลัก "
            "(Normal, Property Crime, Personal Violence) พร้อมเชื่อมโยงกับระบบจดจำใบหน้าความมั่นใจสูงในการคุม Geofence "
            "และวิเคราะห์ระดับความเสี่ยงผ่านโครงข่ายประสาทเทียมส่วนกลาง"
        ),
        
        # 6. วัตถุประสงค์
        "เพื่อตรวจจับวัตถุอันตราย เช่น อาวุธปืนและมีด ด้วยโมเดลที่ฝึกสอนเอง": (
            "-  เพื่อพัฒนาระบบจดจำใบหน้า (Face Recognition) และสืบค้นเปรียบเทียบในคลังข้อมูลบุคคลเฝ้าระวัง (Watchlist 1:N Search) ด้วยความถูกต้องแม่นยำสูง"
        ),
        
        # 7. ขอบเขตของโครงงาน
        "ไม่ครอบคลุมการระบุตัวตนบุคคล (Face Recognition) เพื่อความเป็นส่วนตัว": (
            "-  ครอบคลุมการจดจำใบหน้าและระบุตัวตนของบุคคลใน Watchlist แบบ 1:N Search เพื่อใช้ในระบบตรวจสอบสิทธิ์เข้าออกพื้นที่อัจฉริยะ"
        ),
        "ประเมินความเสี่ยงระดับ LOW / MEDIUM / HIGH / CRITICAL": (
            "-  ประเมินความเสี่ยงระดับ LOW / MEDIUM / HIGH / CRITICAL ผ่านโครงข่ายประสาทเทียม Risk MLP จากข้อมูลพฤติกรรม 6 มิติ"
        ),
        
        # 8. ประโยชน์ที่คาดว่าจะได้รับ
        "ประเมินระดับความเสี่ยงได้อย่างแม่นยำและแจ้งเตือนได้ทันเหตุการณ์": (
            "-  ประเมินระดับความเสี่ยงและคัดกรองบุคคลแปลกหน้า (Stranger) ได้อย่างแม่นยำและแจ้งเตือนได้ทันเหตุการณ์"
        ),
        
        # 9. ทฤษฎีและหลักการ
        "ใช้ YOLOv8 ในการตรวจจับบุคคลและวัตถุในภาพ และติดตามด้วย BoT-SORT": (
            "ใช้ YOLOv8 ในการตรวจจับบุคคลในภาพ และติดตามด้วย ByteTrack ซึ่งมี Kalman Filter และกลไกทำ Association ที่รวดเร็วเพื่อวิเคราะห์ทิศทางการเคลื่อนตัวในเฟรมสดแบบ Real-time"
        ),
        "-  Appearance Classification: โมเดล EfficientNet-B0 ที่ฝึกสอนเองสำหรับวิเคราะห์เครื่องแต่งกายและการปิดบังใบหน้า": (
            "-  Face Recognition: ใช้แบบจำลอง SCRFD ในการระบุพิกัดใบหน้าและ Affine alignment ขนาด 112x112 ร่วมกับ ArcFace (w600k_r50) สกัดรหัสใบหน้า 512 มิติในการสืบค้นเปรียบเทียบผ่านระยะเวกเตอร์แบบ Cosine Similarity"
        ),
        "-  Behavior Recognition: โมเดล LSTM บน Pose Keypoints สำหรับจำแนกพฤติกรรม": (
            "-  Behavior & Action Recognition: โมเดลจำแนกพฤติกรรม 3 คลาส (Normal, Property Crime, Personal Violence) โดยใช้โครงข่ายประสาทเทียม EfficientNet-B0 + Temporal GRU (PyTorch) ฝึกสอนบนฐานข้อมูลบูรณาการ (UCF-Crime, Youtube Robbery, SCVD) พร้อมตัวกรองความมั่นใจ (Confidence Threshold Filter 70%)"
        ),
        "-  Weapon Detection: โมเดล YOLOv8 ที่ฝึกสอนเองจาก Dataset จริง เพื่อตรวจจับอาวุธปืนและมีด": (
            "-  Risk Assessment: ประเมินความเสี่ยงผ่านโครงข่ายประสาทเทียม Feed-forward MLP (PyTorch) ที่รับฟีเจอร์พฤติกรรม 6 มิติมาทำนายระดับความเสี่ยงและค่าความมั่นใจแบบไดนามิก"
        ),
        
        # 10. ขั้นตอนดำเนินการ
        "รวบรวมและเตรียม Dataset สำหรับฝึกสอนโมเดล Weapon Detection และ Appearance Classification": (
            "-  รวบรวมและเตรียม Dataset สำหรับฝึกสอนโมเดล Action Classifier และสร้างฐานข้อมูลใบหน้า Watchlist Database"
        ),
        "พัฒนาโมดูลตรวจจับและติดตามบุคคลด้วย YOLOv8 และ BoT-SORT ร่วมกับ Kalman Filter": (
            "-  พัฒนาโมดูลตรวจจับและติดตามบุคคลด้วย YOLOv8 และ ByteTrack ร่วมกับ Kalman Filter"
        ),
        "พัฒนาและฝึกสอนโมเดล EfficientNet-B0, YOLOv8-Pose + LSTM และ YOLOv8n สำหรับ Weapon Detection": (
            "-  พัฒนาและฝึกสอนโมเดลจำแนกพฤติกรรม (EfficientNet-B0 + GRU), จดจำใบหน้า (SCRFD + ArcFace) และ MLP สำหรับวิเคราะห์ความเสี่ยง"
        ),
        "พัฒนาโมดูลประเมินความเสี่ยงโดยผสมผสาน Rule-based และผลลัพธ์จากโมเดล AI ทั้งหมด": (
            "-  พัฒนาโมดูลประเมินความเสี่ยงโดยผสมผสานโครงข่ายประสาทเทียม MLP และลอจิก Geofencing"
        ),
        
        # 11. สรุป
        "ประเมินความเสี่ยงจากพฤติกรรมและลักษณะภายนอก ตลอดจนตรวจจับวัตถุอันตราย": (
            "ประเมินความเสี่ยงจากพฤติกรรมโดยใช้โมเดลจำแนกพฤติกรรม 3 คลาส และการระบุพิกัดสิทธิ์ใบหน้าผ่านระบบ Face Recognition"
        )
    }
    
    for p in doc.paragraphs:
        p_text = p.text
        for search_str, replacement in replacements.items():
            if search_str in p_text:
                print(f"  Replacing in paragraph: '{search_str[:30]}...' -> '{replacement[:30]}...'")
                p.text = replacement
                break # target only one replacement per paragraph loop
                
    # ── 2. การแก้ไขข้อความในตาราง (Tables) ───────────────────────────
    print("\nUpdating tables...")
    
    # --- ตารางที่ 1: ตารางสรุปเทคโนโลยี ---
    table_tech = None
    for t in doc.tables:
        if len(t.rows) > 1 and "Object Detection" in t.rows[1].cells[0].text:
            table_tech = t
            break
            
    if table_tech:
        print("  Found Technology Table. Updating rows...")
        # กำหนดข้อมูลแถวใหม่สำหรับตารางเทคโนโลยี
        new_tech_data = [
            ("Object Detection & Tracking", "YOLOv8 + ByteTrack (Ultralytics)"),
            ("Face Detection & Alignment", "SCRFD (InsightFace)"),
            ("Face Embedding", "ArcFace (ONNX w600k_r50)"),
            ("Action Classifier", "EfficientNet-B0 + GRU (PyTorch)"),
            ("Risk Classifier", "Feed-forward MLP (PyTorch)"),
            ("Image Processing", "OpenCV + ONNX Runtime")
        ]
        
        # ลบแถวเก่าและสร้างแถวใหม่เพื่อความปลอดภัยในการจัดหน้าตาราง
        # แถวแรก (Header) เก็บไว้
        while len(table_tech.rows) > 1:
            # ลบแถวที่สองเรื่อยๆ จนเหลือแต่ Header
            tbl = table_tech._tbl
            tbl.remove(table_tech.rows[1]._tr)
            
        # เติมแถวใหม่เข้าไปในตาราง
        for area, tech in new_tech_data:
            row_cells = table_tech.add_row().cells
            row_cells[0].text = area
            row_cells[1].text = tech
            
    # --- ตารางที่ 2: ตาราง Gantt Chart (แผนการดำเนินงาน) ---
    table_gantt = None
    for t in doc.tables:
        if len(t.rows) > 1 and "ศึกษาและทบทวนเทคโนโลยี" in t.rows[1].cells[1].text:
            table_gantt = t
            break
            
    if table_gantt:
        print("  Found Gantt Chart Table. Updating row texts...")
        # อัปเดตรายการแถว Gantt Chart
        table_gantt.rows[1].cells[1].text = "ศึกษาและทบทวนเทคโนโลยี YOLOv8, ByteTrack, EfficientNet-B0, GRU, SCRFD, ArcFace"
        table_gantt.rows[3].cells[1].text = "พัฒนาโมดูลตรวจจับและติดตามบุคคล (YOLOv8 + ByteTrack)"
        table_gantt.rows[4].cells[1].text = "พัฒนาและประเมินระบบจดจำใบหน้า (SCRFD + ArcFace) และจำแนกพฤติกรรม (EfficientNet-B0 + GRU)"
        table_gantt.rows[5].cells[1].text = "พัฒนาโมดูลประเมินความเสี่ยงด้วยโครงข่ายประสาทเทียม MLP"
        
    doc.save(updated_path)
    print(f"\nSuccessfully saved updated document to: {updated_path}")

if __name__ == "__main__":
    main()
